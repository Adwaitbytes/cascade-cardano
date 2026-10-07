"""Assembles demo/out/cascade-pitch.pptx and .pdf from the rendered slide PNGs.

Every slide is one full-bleed image with its spoken pitch in the speaker notes. The demo slide
embeds demo/out/cascade-demo.mp4 over its poster when the file exists; otherwise it keeps the poster.

Run from the repo root: uv run --with python-pptx --with pillow python demo/deck2/build.py
"""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

DECK = Path(__file__).resolve().parent
OUT = DECK.parent / "out"
VIDEO = Path(os.environ.get("CASCADE_DEMO_VIDEO", OUT / "cascade-demo.mp4"))
VIDEO_SLIDE = 7
# The video frame on slide 7, in 1920x1080 slide pixels (deck.html #videoFrame).
VIDEO_BOX = (310, 320, 1300, 731)
SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)


def px(value: int, total_px: int, total_emu: int) -> Emu:
    return Emu(round(value * total_emu / total_px))


def main() -> int:
    pngs = sorted((DECK / "png").glob("slide-*.png"))
    notes = json.loads((DECK / "notes.json").read_text())
    if len(pngs) != len(notes):
        print(f"build: {len(pngs)} slides but {len(notes)} notes", file=sys.stderr)
        return 1

    OUT.mkdir(exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]
    embedded = False

    for index, (png, note) in enumerate(zip(pngs, notes), start=1):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, SLIDE_W, SLIDE_H)
        slide.notes_slide.notes_text_frame.text = note
        if index == VIDEO_SLIDE and VIDEO.exists():
            x, y, w, h = VIDEO_BOX
            with Image.open(png) as full:
                scale = full.width / 1920
                poster = full.crop((round(x * scale), round(y * scale), round((x + w) * scale), round((y + h) * scale)))
                poster_path = DECK / "png" / "video-poster.png"
                poster.save(poster_path)
            slide.shapes.add_movie(
                str(VIDEO),
                px(x, 1920, SLIDE_W), px(y, 1080, SLIDE_H), px(w, 1920, SLIDE_W), px(h, 1080, SLIDE_H),
                poster_frame_image=str(poster_path),
                mime_type="video/mp4",
            )
            embedded = True

    pptx_path = OUT / "cascade-pitch.pptx"
    prs.save(pptx_path)

    pages = [Image.open(p).convert("RGB").resize((1920, 1080), Image.LANCZOS) for p in pngs]
    pdf_path = OUT / "cascade-pitch.pdf"
    pages[0].save(pdf_path, save_all=True, append_images=pages[1:], resolution=144)

    print(f"wrote {pptx_path} ({len(pngs)} slides, video {'embedded' if embedded else 'not found, poster only'})")
    print(f"wrote {pdf_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
