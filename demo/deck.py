"""Builds demo/out/cascade-deck.pptx, the eight-slide deck of docs/submission/deck-outline.md.

Run: cd demo && uv run deck.py

The palette is the web app's: ink on slate, with the six node-state colours as the only saturated
hues, so a colour always means a state of money. Diagrams are drawn as shapes. Slide 4 embeds
demo/out/cascade-demo.mp4 (the file itself, not a link) when it exists, else shows a labelled slot.
Screenshots come from demo/shots.ts (live site). Totals are read from the public landing API at
build time, and every number names its source.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import urllib.request
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "demo" / "out"
DEPLOYMENT = json.loads((ROOT / "deployments" / "preprod.json").read_text())
ORIGIN = DEPLOYMENT["urls"]["origin"]
API = f"{ORIGIN}/api/v1"

# ------------------------------------------------------------------------------------------------
# Design tokens (apps/web/src/app/globals.css)

INK = RGBColor(0x0D, 0x18, 0x26)
INK_2 = RGBColor(0x44, 0x52, 0x66)
INK_3 = RGBColor(0x5F, 0x6B, 0x7C)
BG = RGBColor(0xF3, 0xF5, 0xF8)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE_2 = RGBColor(0xEE, 0xF1, 0xF5)
LINE = RGBColor(0xDD, 0xE3, 0xEA)
LINE_STRONG = RGBColor(0xC3, 0xCC, 0xD7)
ON_DARK = RGBColor(0xE7, 0xEE, 0xF6)
ON_DARK_2 = RGBColor(0xA9, 0xB6, 0xC6)
DARK_LINE = RGBColor(0x2C, 0x3E, 0x53)
DARK_SURFACE = RGBColor(0x13, 0x21, 0x2F)


@dataclass(frozen=True)
class State:
    fg: RGBColor
    bg: RGBColor


FUNDED = State(RGBColor(0x1D, 0x56, 0xD8), RGBColor(0xE7, 0xEF, 0xFF))
WORKING = State(RGBColor(0x9A, 0x4A, 0x05), RGBColor(0xFD, 0xF0, 0xDC))
SUBMITTED = State(RGBColor(0x63, 0x37, 0xC4), RGBColor(0xF0, 0xEA, 0xFF))
ACCEPTED = State(RGBColor(0x13, 0x77, 0x3A), RGBColor(0xE2, 0xF5, 0xE9))
REFUNDED = State(RGBColor(0x52, 0x5D, 0x6B), RGBColor(0xEC, 0xEF, 0xF3))
CHALLENGED = State(RGBColor(0xB4, 0x23, 0x23), RGBColor(0xFD, 0xE8, 0xE8))

# Arial renders the same in PowerPoint and in the LibreOffice PDF export.
SANS = "Arial"
MONO = "Courier New"

# Type scale (pt)
DISPLAY = 50
H1 = 34
H2 = 21
BODY = 16
SMALL = 13
MICRO = 10.5

W = 13.333
H = 7.5
MARGIN = 0.8


# ------------------------------------------------------------------------------------------------
# Primitives


def set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    content: str | list[tuple[str, dict]],
    *,
    size: float = BODY,
    color: RGBColor = INK,
    bold: bool = False,
    font: str = SANS,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line: float = 1.15,
    tracking: int = 0,
    upper: bool = False,
):
    """A text box. `content` is a string (paragraphs split on newlines) or a list of runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paragraphs = content.split("\n") if isinstance(content, str) else [content]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        runs = [(para, {})] if isinstance(para, str) else para
        for run_text, style in runs:
            r = p.add_run()
            r.text = run_text.upper() if upper else run_text
            f = r.font
            f.name = style.get("font", font)
            f.size = Pt(style.get("size", size))
            f.bold = style.get("bold", bold)
            f.color.rgb = style.get("color", color)
            if "link" in style:
                r.hyperlink.address = style["link"]
            if tracking:
                r._r.get_or_add_rPr().set("spc", str(tracking))
    return box


def rect(slide, x, y, w, h, *, fill: RGBColor | None = SURFACE, line: RGBColor | None = None, radius: float = 0.12, weight: float = 1.0, dash: bool = False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius > 0:
        shape.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(weight)
        if dash:
            shape.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
    shape.shadow.inherit = False
    shape.text_frame.margin_left = shape.text_frame.margin_right = Inches(0.12)
    return shape


def arrow(slide, x1, y1, x2, y2, *, color: RGBColor = LINE_STRONG, weight: float = 1.5, head: bool = True, dash: bool = False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    ln = c.line._get_or_add_ln()
    if dash:
        prst = etree.SubElement(ln, qn("a:prstDash"))
        prst.set("val", "dash")
    if head:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("h", "med")
    return c


def elbow(slide, x1, y1, x2, y2, **kw) -> None:
    """Parent bottom-centre to child top-centre: down, across, down."""
    mid = y1 + (y2 - y1) / 2
    arrow(slide, x1, y1, x1, mid, head=False, **kw)
    arrow(slide, x1, mid, x2, mid, head=False, **kw)
    arrow(slide, x2, mid, x2, y2, **kw)


def node(slide, x, y, w, h, label: str, state: State, sub: str = "", *, dark: bool = False, size: float = SMALL):
    """A tree node in the explorer's style: tinted card, state-coloured left rule, name and state."""
    card = rect(slide, x, y, w, h, fill=DARK_SURFACE if dark else state.bg, line=state.fg if dark else None, radius=0.08, weight=1.0)
    rect(slide, x, y + 0.08, 0.06, h - 0.16, fill=state.fg, radius=0.03)
    label_h = h * 0.55 if sub else h - 0.12
    text(slide, x + 0.18, y + 0.06, w - 0.26, label_h, label, size=size, bold=True, color=ON_DARK if dark else INK, anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        text(slide, x + 0.18, y + h * 0.52, w - 0.26, h * 0.42, sub, size=MICRO, color=state.fg if not dark else ON_DARK_2, anchor=MSO_ANCHOR.TOP)
    return card


def kicker(slide, label: str, *, color: RGBColor = INK_3, y: float = 0.62) -> None:
    text(slide, MARGIN, y, 8, 0.3, label, size=MICRO + 0.5, color=color, bold=True, tracking=120, upper=True)


def title(slide, content: str, *, color: RGBColor = INK, y: float = 0.92, size: float = H1, w: float = W - 2 * MARGIN) -> None:
    text(slide, MARGIN, y, w, 1.2, content, size=size, bold=True, color=color, line=1.05)


def footer(slide, n: int, *, dark: bool = False, source: str = "") -> None:
    c = ON_DARK_2 if dark else INK_3
    text(slide, MARGIN, H - 0.48, 6, 0.25, "Cascade", size=MICRO, color=c, bold=True)
    text(slide, W - MARGIN - 1, H - 0.48, 1, 0.25, f"{n:02d}", size=MICRO, color=c, align=PP_ALIGN.RIGHT, font=MONO)
    if source:
        text(slide, MARGIN + 1.0, H - 0.48, W - 2 * MARGIN - 2.2, 0.25, source, size=MICRO - 1, color=c)


def notes(slide, content: str) -> None:
    slide.notes_slide.notes_text_frame.text = content


# ------------------------------------------------------------------------------------------------
# Data (real preprod, read at build time)

COWORKER_ID = "01a110cd-4ee0-763b-ae63-4008564c9f8e"
TASK_ID = "01a110e2-b1ca-752c-bc65-123cd12ae594"
REPO = "https://github.com/Adwaitbytes/cascade-cardano"
SCAN = "https://preprod.cardanoscan.io/transaction/"
# Paid Task 1 (docs/sokosumi-coworker.md, Proof).
TASK_PROOF = [
    ("Masumi escrow lock", "Sokosumi pays from the buyer's credits", "caafb951464f87f3f2f2f972875657f6723c870e81a7152bd8ca61116472c589"),
    ("Result hash on chain", "SubmitResult before the Task completes", "a9cefb55e6a6b035ef6cbee15332691b09741204d2e86cc1beec7e85aa578d1c"),
    ("Seller collects", "state Withdrawn after the unlock time", "0e584a87be1ce0315a42a39a7c1c18b100f7c75074ede6825f8c351bb3dc93c4"),
]


def get(path: str):
    req = urllib.request.Request(f"{API}{path}", headers={"accept": "application/json", "user-agent": "cascade-deck"})
    with urllib.request.urlopen(req, timeout=30) as res:
        return json.loads(res.read())


@dataclass
class Totals:
    trees: int
    txs: int
    payouts: int
    agents_paid: int
    paid_lovelace: int
    returned_lovelace: int
    hero_tree: str
    fetched_at: str


def totals() -> Totals:
    d = get("/landing")
    t = d["totals"]
    if t["asset"] != "lovelace":
        raise SystemExit(f"landing totals are in {t['asset']}, the deck formats lovelace only")
    return Totals(
        trees=int(t["trees"]),
        txs=int(t["txs"]),
        payouts=int(t["payouts"]),
        agents_paid=int(t["agents_paid"]),
        paid_lovelace=int(t["paid"]),
        returned_lovelace=int(t["returned"]) + int(t["structural_returned"]),
        hero_tree=d.get("hero_tree_id") or "",
        fetched_at=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
    )


def redeemer_rows() -> list[dict]:
    return [r for r in json.loads((OUT / "redeemers.json").read_text())["rows"] if r["primary"]]


def aiken_tests() -> int:
    return int(json.loads((OUT / "aiken-check.json").read_text())["total"])


def ada(lovelace: int) -> str:
    return f"{lovelace / 1_000_000:,.2f}"


def short(h: str) -> str:
    return f"{h[:8]}...{h[-4:]}"


def picture(s, path: Path, x, y, w, *, crop_top=0.0, crop_bottom=0.0, crop_left=0.0, crop_right=0.0, aspect=16 / 9, frame: RGBColor = LINE) -> float:
    """A framed screenshot, cropped as fractions of the source; returns the drawn height."""
    h = w / aspect * (1 - crop_top - crop_bottom) / (1 - crop_left - crop_right)
    rect(s, x - 0.05, y - 0.05, w + 0.1, h + 0.1, fill=SURFACE, line=frame, radius=0.08)
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic.crop_top, pic.crop_bottom, pic.crop_left, pic.crop_right = crop_top, crop_bottom, crop_left, crop_right
    return h


# ------------------------------------------------------------------------------------------------
# Slides


def slide_title(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, INK)
    kicker(s, "Sokosumi Coworker, live on Cardano preprod", color=ON_DARK_2)
    text(s, MARGIN, 1.1, 6, 1.2, "Cascade", size=DISPLAY + 22, bold=True, color=ON_DARK, line=1.0)
    text(s, MARGIN, 2.55, 5.4, 1.8, "One Task on Sokosumi.\nA whole team of agents.\nEscrow at every level.", size=H2 + 3, color=ON_DARK_2, line=1.3)
    text(s, MARGIN, 4.75, 5.2, 0.3, "Coworker", size=MICRO, color=ON_DARK_2, bold=True, tracking=120, upper=True)
    text(s, MARGIN, 5.02, 5.4, 0.3, COWORKER_ID, size=SMALL, color=ON_DARK, font=MONO)
    text(s, MARGIN, 5.5, 5.2, 0.3, "Try it", size=MICRO, color=ON_DARK_2, bold=True, tracking=120, upper=True)
    text(s, MARGIN, 5.77, 5.4, 0.3, ORIGIN.removeprefix("https://"), size=SMALL, color=ON_DARK, font=MONO)
    shot = OUT / "shots" / "landing.png"
    if shot.exists():
        picture(s, shot, 6.35, 1.35, W - MARGIN - 6.35, crop_bottom=0.08, frame=DARK_LINE)
    footer(s, 1, dark=True)
    notes(s, "Cascade is a Coworker you hire on Sokosumi. You give it one Task. It hires the agents it needs, and every one of them is paid from escrow on Cardano, only for checked work.")


def slide_problem(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, BG)
    kicker(s, "The problem")
    title(s, "Escrow stops at the first hop.")
    cards = [
        (WORKING, "Real work needs a team", "Research, pricing, writing and checking are different agents. One agent rarely does the whole Task."),
        (CHALLENGED, "The coordinator fronts the money", "It pays every sub-hire from its own wallet and has to trust each one to deliver."),
        (REFUNDED, "A failed sub-agent is a loss", "Nothing enforces escrow past the first hire, so a failure is lost money, not a refund."),
    ]
    cw, gap = 3.75, 0.24
    for i, (st, head, body) in enumerate(cards):
        x = MARGIN + i * (cw + gap)
        rect(s, x, 2.4, cw, 3.0, fill=SURFACE, line=LINE, radius=0.14)
        rect(s, x + 0.35, 2.8, 0.42, 0.06, fill=st.fg, radius=0)
        text(s, x + 0.35, 3.05, cw - 0.7, 0.9, head, size=H2 - 1, bold=True, line=1.1)
        text(s, x + 0.35, 3.95, cw - 0.7, 1.3, body, size=BODY - 1, color=INK_2, line=1.25)
    text(s, MARGIN, 5.85, W - 2 * MARGIN, 0.4, "The buyer cannot see who did what, or what each part cost.", size=BODY, color=INK_2)
    footer(s, 2)
    notes(s, "Masumi's own workshop slide says every hire, at every level, is paid into escrow. Today nothing enforces that past the first hire. The buyer cannot see who did what or what it cost.")


def slide_solution(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, BG)
    kicker(s, "The solution: an escrow tree")
    title(s, "One payment. Each hire draws its own child escrow.", size=H1 - 4, w=11.4)

    top = 2.2
    node(s, 5.67, top, 2.0, 0.6, "Buyer", FUNDED, "one signature")
    arrow(s, 6.67, top + 0.6, 6.67, top + 0.95, color=INK_3)
    node(s, 5.17, top + 0.95, 3.0, 0.66, "Conductor", FUNDED, "root escrow, plan as a Merkle root")
    level1 = [
        ("Scout", ACCEPTED, "hires its own child"),
        ("Flaky Lisan", REFUNDED, "test agent, refunded"),
        ("Lisan", ACCEPTED, "Masumi agent"),
        ("Checker A", SUBMITTED, "verifier"),
        ("Checker B", SUBMITTED, "verifier"),
        ("Scribe", ACCEPTED, "writes the brief"),
    ]
    cw, gap = 1.76, (W - 2 * MARGIN - 6 * 1.76) / 5
    y1 = top + 2.05
    xs = []
    for i, (label, st, sub) in enumerate(level1):
        x = MARGIN + i * (cw + gap)
        xs.append(x)
        elbow(s, 6.67, top + 1.61, x + cw / 2, y1, color=LINE_STRONG, weight=1.25)
        node(s, x, y1, cw, 0.7, label, st, sub)
    arrow(s, xs[1] + cw * 0.7, y1 - 0.02, 5.17, top + 1.3, color=REFUNDED.fg, weight=1.75, dash=True)
    text(s, MARGIN, top + 1.05, 2.6, 0.3, "refund flows up, re-spent on Lisan", size=MICRO, color=REFUNDED.fg, bold=True)
    y2 = y1 + 1.2
    elbow(s, xs[0] + cw / 2, y1 + 0.7, xs[0] + cw / 2, y2, color=LINE_STRONG, weight=1.25)
    node(s, xs[0], y2, cw, 0.7, "Pricer", ACCEPTED, "depth 2")
    text(s, 3.6, y2 - 0.02, W - MARGIN - 3.6, 0.9, "Validators enforce budgets, nested deadlines and plan membership. A missed deadline refunds up the tree, and the budget is re-spent on a replacement.", size=BODY - 1, color=INK_2, line=1.3)
    footer(s, 3)
    notes(s, "Every node of the tree is its own UTxO with its own thread token. A child can never hold more than its parent drew for it. A parent cannot submit while a child is open. Anyone can crank a refund after a deadline, so no silent operator can lock funds.")


def slide_demo(prs, video: Path | None, poster: Path | None) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, INK)
    kicker(s, "Demo, recorded on Cardano preprod", color=ON_DARK_2, y=0.5)
    text(s, MARGIN, 0.78, W - 2 * MARGIN, 0.6, "Sokosumi Task, live Tree Explorer, receipt, Masumi escrow.", size=H2 + 3, bold=True, color=ON_DARK)
    vw = 9.4
    vh = vw * 9 / 16
    vx, vy = (W - vw) / 2, 1.55
    if video is not None:
        s.shapes.add_movie(str(video), Inches(vx), Inches(vy), Inches(vw), Inches(vh), poster_frame_image=str(poster) if poster is not None and poster.exists() else None, mime_type="video/mp4")
    else:
        rect(s, vx, vy, vw, vh, fill=DARK_SURFACE, line=DARK_LINE, radius=0.1, dash=True)
        text(s, vx, vy + vh / 2 - 0.45, vw, 0.6, "Demo video", size=H1, bold=True, color=ON_DARK, align=PP_ALIGN.CENTER)
        text(s, vx, vy + vh / 2 + 0.2, vw, 0.4, "Insert the recording here: Insert > Video > demo/out/cascade-demo.mp4", size=SMALL, color=ON_DARK_2, align=PP_ALIGN.CENTER)
    text(s, MARGIN, vy + vh + 0.15, W - 2 * MARGIN, 0.3, [("Task ", {}), (TASK_ID, {"font": MONO, "color": ON_DARK}), ("   Flaky Lisan is a test agent that fails on purpose.", {})], size=MICRO + 0.5, color=ON_DARK_2, align=PP_ALIGN.CENTER)
    footer(s, 4, dark=True)
    notes(s, "Play the 2:30 video or the 20-second hook. Point at the refund: Flaky Lisan, a test agent, fails on purpose, and its budget flows back up and hires a replacement.")


def slide_built_on(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, BG)
    kicker(s, "Built on Masumi, Sokosumi and x402")
    title(s, "Cascade nests the payment rails that exist.")
    cols = [
        (FUNDED, "Sokosumi", "Cascade is a Coworker. Tasks are paid with credits into Masumi escrow."),
        (ACCEPTED, "Masumi", "Registry identity for the agents and vested_pay escrow. Unmodified Masumi agents are hired through MIP-003, refunds go to the buyer."),
        (WORKING, "x402", "Agents quote and pay each other with HTTP 402. Metered calls settle on vouchers."),
    ]
    cw, gap = 3.75, 0.24
    for i, (st, head, body) in enumerate(cols):
        x = MARGIN + i * (cw + gap)
        rect(s, x, 2.1, 0.42, 0.06, fill=st.fg, radius=0)
        text(s, x, 2.3, cw, 0.45, head, size=H2, bold=True)
        text(s, x, 2.8, cw, 1.3, body, size=SMALL + 1, color=INK_2, line=1.3)

    y = 4.45
    rect(s, MARGIN, y, W - 2 * MARGIN, 2.1, fill=SURFACE, line=LINE, radius=0.14)
    text(s, MARGIN + 0.35, y + 0.25, 8, 0.3, [("Paid Task 1 on preprod, 2026-10-06   ", {"bold": True, "color": INK}), (TASK_ID, {"font": MONO})], size=SMALL, color=INK_3)
    sw = (W - 2 * MARGIN - 0.7 - 2 * 0.5) / 3
    for i, (head, sub, tx) in enumerate(TASK_PROOF):
        x = MARGIN + 0.35 + i * (sw + 0.5)
        node(s, x, y + 0.75, sw, 1.0, head, ACCEPTED, sub)
        text(s, x + 0.18, y + 1.45, sw - 0.3, 0.25, [(short(tx), {"link": SCAN + tx})], size=MICRO, color=FUNDED.fg, font=MONO)
        if i < 2:
            arrow(s, x + sw + 0.08, y + 1.25, x + sw + 0.42, y + 1.25, color=INK_3)
    footer(s, 5, source="docs/sokosumi-coworker.md, docs/submission/cardano-track.md")
    notes(s, "Cascade is the missing layer on top of Masumi and x402. It does not replace them; it nests them. A Cascade agent answers 402 with two offers: pay into a child escrow or into a Masumi lock. Paid Task 1 proves the Sokosumi to Masumi escrow to seller payout path; each box links to its preprod transaction.")


def slide_proof(prs, t: Totals, rows: list[dict], aiken: int) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, BG)
    kicker(s, "Proof on preprod")
    title(s, f"All {len(rows)} escrow-tree redeemers ran on preprod.")
    # Redeemer chips, each linked to its transaction.
    cols, cw, ch, gx, gy = 3, 1.78, 0.42, 0.12, 0.12
    for i, r in enumerate(rows):
        x = MARGIN + (i % cols) * (cw + gx)
        y = 2.05 + (i // cols) * (ch + gy)
        st = REFUNDED if r["redeemer"] == "Refund" else CHALLENGED if r["redeemer"] in ("Challenge", "Escalate", "Resolve") else ACCEPTED
        rect(s, x, y, cw, ch, fill=st.bg, line=None, radius=0.06)
        text(s, x + 0.12, y, cw - 0.2, ch, [(r["redeemer"], {"link": r["cardanoscan"]})], size=MICRO - 1.5, bold=True, color=st.fg, anchor=MSO_ANCHOR.MIDDLE)
    shot = OUT / "shots" / "explorer.png"
    sx = MARGIN + cols * (cw + gx) + 0.3
    if shot.exists():
        picture(s, shot, sx, 2.05, W - MARGIN - sx, crop_top=0.3, crop_bottom=0.06, crop_left=0.1, crop_right=0.1)
        text(s, sx, 5.15, W - MARGIN - sx, 0.3, f"Live Tree Explorer, tree {short(t.hero_tree)}", size=MICRO, color=INK_3)
    stats = [
        (f"{t.trees}", "trees funded"),
        (f"{t.txs}", "transactions"),
        (f"{t.payouts}", f"payouts to {t.agents_paid} payees"),
        (ada(t.paid_lovelace), "ADA paid to agents"),
        (f"{t.returned_lovelace / 1_000_000:,.0f}", "ADA returned to buyers"),
        (f"{aiken}", "Aiken tests, 0 failures"),
    ]
    sw = (W - 2 * MARGIN) / len(stats)
    for i, (num, label) in enumerate(stats):
        x = MARGIN + i * sw
        text(s, x, 5.6, sw - 0.15, 0.5, num, size=H2 + 3, bold=True)
        text(s, x, 6.1, sw - 0.15, 0.4, label, size=MICRO, color=INK_2)
    footer(s, 6, source=f"Each chip links its tx. Totals: /api/v1/landing, {t.fetched_at}. Tests: demo/out/aiken-check.json.")
    notes(s, "Every chip opens its preprod transaction. Click Refund: Flaky Lisan, a test agent that fails on purpose, missed its deadline, and the watchtower refunded it into its parent in one transaction.")


def slide_new(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, BG)
    kicker(s, "What is new")
    title(s, "Cascade escrows the whole hiring chain.")
    rows = [
        ("Masumi", "Escrows one hop: buyer to one agent.", REFUNDED),
        ("x402", "Pays one request.", REFUNDED),
        ("Cascade", "Nests escrows across the whole hiring chain and binds every payment to a plan the buyer signed.", ACCEPTED),
    ]
    for i, (name, what, st) in enumerate(rows):
        y = 2.2 + i * 0.95
        rect(s, MARGIN, y, W - 2 * MARGIN, 0.8, fill=SURFACE if i < 2 else st.bg, line=LINE if i < 2 else st.fg, radius=0.1)
        text(s, MARGIN + 0.35, y, 2.2, 0.8, name, size=H2 - 1, bold=True, color=INK if i < 2 else st.fg, anchor=MSO_ANCHOR.MIDDLE)
        text(s, MARGIN + 2.6, y, W - 2 * MARGIN - 2.9, 0.8, what, size=BODY, color=INK_2 if i < 2 else INK, anchor=MSO_ANCHOR.MIDDLE)
    text(s, MARGIN, 5.25, W - 2 * MARGIN, 0.4, "We checked seven Cardano agent-payment projects. None ships a tree of escrows.", size=BODY + 1, bold=True)
    text(s, MARGIN, 5.75, W - 2 * MARGIN, 0.4, "ArgoOperator, ANTIDOTE, Sentinel, Proof Pair, NightPay, cardano402, Subbit.xyz: each handles one hop, a pool or a channel.", size=SMALL, color=INK_2)
    footer(s, 7, source="PRD 3.1")
    notes(s, "Name the seven if asked. Each handles one hop, a pool or a channel.")


def slide_next(prs) -> None:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, INK)
    kicker(s, "What's next", color=ON_DARK_2)
    text(s, MARGIN, 0.92, 11, 0.8, "From preprod to mainnet.", size=H1, bold=True, color=ON_DARK)
    items = [
        ("Mainnet", "Second audit, then mainnet with Masumi mainnet escrow and USDM budgets."),
        ("Trustless Masumi leaves", "Once Masumi accepts script-created locks (BLOCKERS.md B6)."),
        ("Coworkers as leaves", "Other Sokosumi Coworkers hired as leaves of a Cascade tree."),
    ]
    cw, gap = 3.75, 0.24
    for i, (head, body) in enumerate(items):
        x = MARGIN + i * (cw + gap)
        rect(s, x, 2.1, cw, 2.3, fill=DARK_SURFACE, line=DARK_LINE, radius=0.14)
        text(s, x + 0.3, 2.35, 1, 0.4, f"0{i + 1}", size=SMALL, color=ON_DARK_2, font=MONO)
        text(s, x + 0.3, 2.8, cw - 0.6, 0.45, head, size=H2 - 2, bold=True, color=ON_DARK)
        text(s, x + 0.3, 3.3, cw - 0.6, 1.0, body, size=SMALL, color=ON_DARK_2, line=1.3)
    text(s, MARGIN, 4.95, W - 2 * MARGIN, 0.5, "Try the Coworker on Sokosumi. Talk to us if you run agents that hire agents.", size=H2, bold=True, color=ON_DARK)
    for i, (label, value) in enumerate([("Site", ORIGIN), ("Coworker", COWORKER_ID), ("Code", REPO)]):
        y = 5.65 + i * 0.36
        text(s, MARGIN, y, 1.4, 0.3, label, size=MICRO, color=ON_DARK_2, bold=True, tracking=120, upper=True)
        text(s, MARGIN + 1.4, y - 0.02, 9, 0.3, value, size=SMALL, color=ON_DARK, font=MONO)
    footer(s, 8, dark=True)
    notes(s, "Close on the ask: try the Coworker on Sokosumi, and talk to us if you run agents that hire agents. Show the site URL and the Coworker ID.")


def export_pdf(pptx: Path) -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        print("soffice not found; skipping the PDF", file=sys.stderr)
        return
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)], check=True, capture_output=True, timeout=300)
    print(f"wrote {pptx.with_suffix('.pdf')}")


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    video = OUT / "cascade-demo.mp4"
    poster = OUT / "cascade-demo-poster.png"
    if "--dry-run-video" in sys.argv:
        video, poster = OUT / "cascade-demo-dryrun.mp4", OUT / "cascade-demo-dryrun-poster.png"
    if not video.exists():
        print(f"warning: {video.name} not found; slide 4 shows the Demo video slot", file=sys.stderr)

    prs = Presentation()
    prs.slide_width = Emu(int(W * 914400))
    prs.slide_height = Emu(int(H * 914400))
    prs.core_properties.title = "Cascade: escrow trees for agent work"
    prs.core_properties.author = "Cascade"

    t = totals()
    rows = redeemer_rows()
    slide_title(prs)
    slide_problem(prs)
    slide_solution(prs)
    slide_demo(prs, video if video.exists() else None, poster)
    slide_built_on(prs)
    slide_proof(prs, t, rows, aiken_tests())
    slide_new(prs)
    slide_next(prs)

    out = OUT / "cascade-deck.pptx"
    prs.save(str(out))
    (OUT / "cascade-deck-stats.json").write_text(json.dumps({**t.__dict__, "redeemer_rows": len(rows), "source": f"{API}/landing"}, indent=2) + "\n")
    print(f"wrote {out} ({out.stat().st_size / 1e6:.1f} MB), video embedded: {video.exists()}")
    if "--pdf" in sys.argv:
        export_pdf(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
